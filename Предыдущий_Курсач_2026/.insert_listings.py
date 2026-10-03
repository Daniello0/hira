from __future__ import annotations

from pathlib import Path
from textwrap import dedent

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

DOC_PATH = Path('/Users/daniel/Documents/Курсач_2026/Курсовая_Демо.docx')

LISTINGS = [
    (
        'Листинг 3.1 – Фрагмент миграции создания таблицы vacancies',
        dedent('''\
        from alembic import op


        revision = "0001_create_vacancies"
        down_revision = None
        branch_labels = None
        depends_on = None


        def upgrade() -> None:
            op.execute("CREATE EXTENSION IF NOT EXISTS vector")

            op.execute(
                """
                CREATE TABLE vacancies (
                    id SERIAL PRIMARY KEY,
                    title VARCHAR(512) NOT NULL,
                    company VARCHAR(512) NOT NULL,
                    salary VARCHAR(256) NOT NULL,
                    payment_frequency VARCHAR(128) NOT NULL,
                    experience VARCHAR(256) NOT NULL,
                    employment VARCHAR(256) NOT NULL,
                    hiring_format VARCHAR(256) NOT NULL,
                    schedule VARCHAR(256) NOT NULL,
                    hours VARCHAR(256) NOT NULL,
                    work_format VARCHAR(256) NOT NULL,
                    skills TEXT NOT NULL,
                    url VARCHAR(2048) NOT NULL,
                    description TEXT NOT NULL,
                    embedding vector(384),
                    CONSTRAINT uq_vacancies_url UNIQUE (url)
                )
                """
            )


        def downgrade() -> None:
            op.execute("DROP TABLE IF EXISTS vacancies")
        ''').strip(),
    ),
    (
        'Листинг 3.2 – Фрагмент модуля сбора ссылок и парсинга вакансий rabota.by',
        dedent('''\
        # features/parser/file_orchestrator.py
        async def orchestrate_parser_pipeline_async(
            pages_to_parse: int = DEFAULT_PAGES_TO_PARSE,
            vacancy_limit: int = DEFAULT_VACANCY_LIMIT,
            links_file: Path = VACANCY_LINKS_FILE,
            details_file: Path = VACANCY_DETAILS_FILE,
            cleaned_file: Path = VACANCIES_FILE,
        ) -> Any:
            if not has_data(links_file):
                parse_vacancy_links(pages_to_parse=pages_to_parse, output_file=links_file)
            else:
                print(f"Использую существующий файл ссылок: {links_file}")

            parse_vacancy_details(
                input_file=links_file,
                output_file=details_file,
                limit=vacancy_limit,
            )
            cleaned_data = process_csv(input_file=details_file, output_file=cleaned_file)
            affected = await _upsert_cleaned_vacancies(cleaned_data)
            print(f"Upsert в PostgreSQL завершён. Затронуто строк: {affected}")
            return cleaned_data


        # features/parser/url_parser.py
        def parse_vacancy_links(
            pages_to_parse: int = DEFAULT_PAGES_TO_PARSE,
            output_file: Path = VACANCY_LINKS_FILE,
            start_page: int = DEFAULT_START_PAGE,
        ) -> list[str]:
            user_agent = build_user_agent()
            all_links = read_existing_links(output_file)
            seen_links = set(all_links)
            pages = range(start_page, start_page + pages_to_parse)
            for page in track_progress(pages, total=pages_to_parse, description="Сбор url"):
                page_links = parse_search_page(page, user_agent)
                if page_links is None:
                    continue
                if not page_links:
                    break
                new_links = [link for link in page_links if link not in seen_links]
                if not new_links:
                    continue
                all_links.extend(new_links)
                seen_links.update(new_links)
                save_links_to_csv(all_links, output_file)
                time.sleep(random.uniform(REQUEST_DELAY_MIN_SECONDS, REQUEST_DELAY_MAX_SECONDS))
            return all_links


        # features/parser/data_parser.py
        def parse_vacancy(url: str, user_agent: str) -> dict[str, str] | None:
            response_text = fetch_vacancy_page(url, user_agent)
            if response_text is None:
                return None
            soup = BeautifulSoup(response_text, DEFAULT_HTML_PARSER)
            return {
                "title": get_text_safe(soup, "vacancy-title"),
                "company": get_text_safe(soup, "vacancy-company-name"),
                "salary": get_text_safe(soup, "vacancy-salary"),
                "payment_frequency": get_text_safe(soup, "compensation-frequency-text"),
                "experience": get_text_safe(soup, "vacancy-experience"),
                "employment": get_text_safe(soup, "common-employment-text"),
                "hiring_format": get_text_safe(soup, "vacancy-hiring-formats"),
                "schedule": get_text_safe(soup, "work-schedule-by-days-text"),
                "hours": get_text_safe(soup, "working-hours-text"),
                "work_format": get_text_safe(soup, "work-formats-text"),
                "skills": parse_skills(soup),
                "url": url,
                "description": parse_description(soup),
            }


        # features/parser/clear_csv.py
        def process_csv(
            input_file: Path = VACANCY_DETAILS_FILE,
            output_file: Path = VACANCIES_FILE,
        ) -> Any:
            output_file.parent.mkdir(parents=True, exist_ok=True)
            try:
                cleaned_data = process_csv_with_pandas(input_file, output_file)
            except ModuleNotFoundError:
                cleaned_data = process_csv_with_csv(input_file, output_file)
            print(f"Очистка завершена! Файл сохранен как: {output_file}")
            return cleaned_data
        ''').strip(),
    ),
    (
        'Листинг 3.3 – Фрагмент функции формирования взвешенного текстового представления вакансии',
        dedent('''\
        _STRUCTURE_FIELDS: tuple[tuple[str, str], ...] = (
            ("Занятость", "employment"),
            ("График", "schedule"),
            ("Часы", "hours"),
            ("Формат работы", "work_format"),
            ("Опыт", "experience"),
            ("Трудоустройство", "hiring_format"),
        )


        def _repeat_lines(label: str, value: str, times: int) -> list[str]:
            if _skip_value(value):
                return []
            line = f"{label}: {_trim(value)}"
            return [line] * times


        def _description_for_embed(raw: str) -> str | None:
            if not VACANCY_EMBED_INCLUDE_DESCRIPTION:
                return None
            t = _trim(raw)
            if not t:
                return None
            cap = VACANCY_EMBED_DESCRIPTION_MAX_CHARS
            if cap is not None and len(t) > cap:
                return t[:cap].rstrip() + "…"
            return t


        def _vacancy_embed_text(row: Any) -> str:
            chunks: list[str] = []

            title = _trim(row.title)
            if title:
                job_line = f"Должность: {title}"
                chunks.extend([job_line] * VACANCY_EMBED_TITLE_REPEATS)

            chunks.extend(_repeat_lines("Компания", row.company, VACANCY_EMBED_COMPANY_REPEATS))
            for label, attr in _STRUCTURE_FIELDS:
                chunks.extend(
                    _repeat_lines(
                        label,
                        getattr(row, attr),
                        VACANCY_EMBED_STRUCTURE_REPEATS,
                    )
                )
            chunks.extend(_repeat_lines("Зарплата", row.salary, 1))
            chunks.extend(_repeat_lines("Выплаты", row.payment_frequency, 1))

            if not _skip_value(row.skills):
                sk = f"Навыки: {_trim(row.skills)}"
                chunks.extend([sk] * VACANCY_EMBED_SKILLS_REPEATS)

            desc = _description_for_embed(row.description)
            if desc:
                chunks.append(f"Описание: {desc}")

            return "\n\n".join(chunks)
        ''').strip(),
    ),
    (
        'Листинг 3.4 – Фрагмент извлечения и валидации структурированных фильтров через LLM',
        dedent('''\
        def _validate_llm_payload(
            payload: dict[str, Any], allowed_values: dict[str, list[str]]
        ) -> list[str]:
            errors: list[str] = []
            required = set(VACANCY_FILTER_KEYS)
            optional = set(VACANCY_OPTIONAL_KEYS)
            payload_keys = set(payload.keys())
            if not required.issubset(payload_keys):
                errors.append("JSON must contain all required filter keys.")

            role_keywords = payload.get("role_keywords", [])
            if not isinstance(role_keywords, list):
                errors.append("Field 'role_keywords' must be an array.")
            elif not all(isinstance(item, str) for item in role_keywords):
                errors.append("Field 'role_keywords' must contain only strings.")

            for key in VACANCY_FILTER_KEYS:
                value = payload.get(key)
                if not isinstance(value, list):
                    errors.append(f"Field '{key}' must be an array.")
                    continue
                allowed_set = set(allowed_values[key])
                for item in value:
                    if not isinstance(item, dict):
                        errors.append(f"Field '{key}' items must be objects.")
                        continue
                    if set(item.keys()) != {"value", "weight"}:
                        errors.append(
                            f"Field '{key}' items must contain only 'value' and 'weight'."
                        )
                        continue
                    raw_filter_value = item.get("value")
                    raw_weight = item.get("weight")
                    if not isinstance(raw_filter_value, str):
                        errors.append(f"Field '{key}' item 'value' must be a string.")
                        continue
                    if not isinstance(raw_weight, int | float):
                        errors.append(f"Field '{key}' item 'weight' must be numeric.")
                        continue
                    if not 0 <= float(raw_weight) <= 1:
                        errors.append(
                            f"Field '{key}' item 'weight' must be in range [0, 1]."
                        )
                        continue
                    if raw_filter_value not in allowed_set:
                        errors.append(
                            f"Field '{key}' has invalid values: {raw_filter_value}."
                        )
            return errors


        async def get_vacancy_filters_from_text_async(user_message: str) -> dict[str, Any]:
            allowed_values = await build_vacancy_filter_allowed_values()
            prompt = _build_system_prompt(allowed_values)
            base_context = VACANCY_FILTERS_CONTEXT_TEMPLATE.format(
                user_message=user_message.strip()
            )
            context = base_context
            errors: list[str] = []
            for _ in range(LLM_VALIDATION_RETRY_ATTEMPTS):
                try:
                    answer = get_llm_answer_service(prompt, context)
                except ValueError as error:
                    errors = [str(error)]
                else:
                    errors = _validate_llm_payload(answer.payload, allowed_values)
                    if not errors:
                        return answer.payload
                context = _build_retry_context(base_context, errors)
            details = "; ".join(errors) if errors else "Unknown validation error."
            raise RuntimeError(f"LLM response validation failed: {details}")
        ''').strip(),
    ),
    (
        'Листинг 3.5 – Фрагмент алгоритма релаксации фильтров',
        dedent('''\
        def _build_field_drop_order(
            weighted_filters: dict[str, list[tuple[str, float]]],
            protected_keys: set[str],
        ) -> list[tuple[str, float]]:
            fields: list[tuple[str, float]] = []
            for key in VACANCY_FILTER_KEYS:
                values = weighted_filters[key]
                if values:
                    field_weight = max(weight for _, weight in values)
                    fields.append((key, field_weight))

            def _effective_weight(item: tuple[str, float]) -> float:
                key, weight = item
                if key in protected_keys:
                    return min(1.0, weight + PROTECTED_RELAX_WEIGHT_PENALTY)
                return weight

            return sorted(
                fields,
                key=lambda item: (_effective_weight(item), item[0]),
            )


        async def _relax_filters(
            weighted_filters: dict[str, list[tuple[str, float]]],
            *,
            value_counts: dict[str, dict[str, int]],
            role_keywords: list[str],
            min_candidates: int,
        ) -> tuple[dict[str, list[str]], dict[str, list[str]], list[FilterRelaxStep], int]:
            active_filters = _clone_non_empty(_plain_filters_from_weighted(weighted_filters))
            weighted_by_key = {key: list(values) for key, values in weighted_filters.items()}
            dropped_filters = _empty_filters()
            relax_steps: list[FilterRelaxStep] = []
            candidate_count = await _count_filtered_candidates(
                active_filters, role_keywords=role_keywords
            )
            if candidate_count >= min_candidates:
                return _with_all_keys(active_filters), dropped_filters, relax_steps, candidate_count

            protected_keys = set(PROTECTED_RELAX_FILTER_KEYS)
            drop_order = _build_field_drop_order(weighted_by_key, protected_keys)
            for step_index, (key, weight) in enumerate(drop_order, start=1):
                removed_values = [value for value, _ in weighted_by_key.get(key, [])]
                if not removed_values:
                    continue
                candidates_before = candidate_count
                dropped_filters[key].extend(removed_values)
                active_filters.pop(key, None)
                candidate_count = await _count_filtered_candidates(
                    active_filters, role_keywords=role_keywords
                )
                removed_total_count = sum(
                    value_counts.get(key, {}).get(value, 0) for value in removed_values
                )
                relax_steps.append(
                    FilterRelaxStep(
                        step=step_index,
                        filter_key=key,
                        filter_value=", ".join(removed_values),
                        removed_weight=weight,
                        value_count_in_db=removed_total_count,
                        candidates_before=candidates_before,
                        candidates_after=candidate_count,
                    )
                )
                if candidate_count >= min_candidates:
                    break
            return _with_all_keys(active_filters), dropped_filters, relax_steps, candidate_count
        ''').strip(),
    ),
    (
        'Листинг 3.6 – Фрагмент алгоритма доменного реранжирования результатов',
        dedent('''\
        def _query_domain_keywords(query: str, role_keywords: list[str]) -> set[str]:
            query_lower = query.lower()
            matched: set[str] = set(role_keywords)
            for keywords in _DOMAIN_KEYWORDS.values():
                if any(keyword in query_lower for keyword in keywords):
                    matched.update(keywords)
            return matched


        def _rank_with_domain_boost(
            user_query: str,
            role_keywords: list[str],
            vacancies: list[SimilaritySearchResult],
        ) -> list[SimilaritySearchResult]:
            keywords = _query_domain_keywords(user_query, role_keywords)
            if not keywords:
                return vacancies
            scored: list[tuple[float, float, SimilaritySearchResult]] = []
            for vacancy in vacancies:
                haystack = f"{vacancy.title} {vacancy.skills}".lower()
                matches = sum(1 for keyword in keywords if keyword in haystack)
                boosted_distance = (
                    vacancy.cosine_distance - DOMAIN_BOOST_STEP * min(matches, 3)
                )
                scored.append((boosted_distance, vacancy.cosine_distance, vacancy))
            scored.sort(key=lambda item: (item[0], item[1]))
            return [item[2] for item in scored]
        ''').strip(),
    ),
]


def set_cell_border(cell):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in('w:tcBorders')
    if tc_borders is None:
        tc_borders = OxmlElement('w:tcBorders')
        tc_pr.append(tc_borders)
    for edge in ('top', 'left', 'bottom', 'right'):
        element = tc_borders.find(qn(f'w:{edge}'))
        if element is None:
            element = OxmlElement(f'w:{edge}')
            tc_borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), '8')
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), '000000')


def set_cell_margins(cell, top=80, start=80, bottom=80, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for name, value in (('top', top), ('start', start), ('bottom', bottom), ('end', end)):
        element = tc_mar.find(qn(f'w:{name}'))
        if element is None:
            element = OxmlElement(f'w:{name}')
            tc_mar.append(element)
        element.set(qn('w:w'), str(value))
        element.set(qn('w:type'), 'dxa')


def set_paragraph_spacing(paragraph, before=0, after=0, line=1.0):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
    fmt.line_spacing = line


def add_code_listing(doc: Document, title: str, code: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.allow_autofit = True
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    set_cell_border(cell)
    set_cell_margins(cell, top=90, start=120, bottom=90, end=120)

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(p, before=0, after=0, line=1.0)
    run = p.add_run(code)
    run.font.name = 'Courier New'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Courier New')
    run.font.size = Pt(12)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(title_p, before=6, after=0, line=1.0)
    title_run = title_p.add_run(title)
    title_run.bold = True
    title_run.font.name = 'Times New Roman'
    title_run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    title_run.font.size = Pt(13)

    note_p = doc.add_paragraph()
    note_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(note_p, before=0, after=6, line=1.0)
    note_run = note_p.add_run('Примечание – Источник: собственная разработка.')
    note_run.font.name = 'Times New Roman'
    note_run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    note_run.font.size = Pt(12)

    doc.add_paragraph()


doc = Document(str(DOC_PATH))
all_text = '\n'.join(p.text for p in doc.paragraphs)
if 'Добавленные листинги' in all_text:
    raise SystemExit('Блок "Добавленные листинги" уже присутствует в документе.')

page_break = doc.add_paragraph()
page_break.add_run().add_break(WD_BREAK.PAGE)

heading = doc.add_paragraph()
heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_paragraph_spacing(heading, before=0, after=8, line=1.0)
heading_run = heading.add_run('Добавленные листинги')
heading_run.bold = True
heading_run.font.name = 'Times New Roman'
heading_run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
heading_run.font.size = Pt(14)

for title, code in LISTINGS:
    add_code_listing(doc, title, code)

doc.save(str(DOC_PATH))
print('Inserted listings:', len(LISTINGS))
print('Updated:', DOC_PATH)
