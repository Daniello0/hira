#!/usr/bin/env python3
"""Append missing appendix listings + fix instructions to Курсовая_Вторая_Демо.docx (end only)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent
DOC_PATH = ROOT / "Курсовая_Вторая_Демо.docx"
PROJECT = Path(__file__).resolve().parents[1] / "ai-job-finder-python"


def _set_run_font(run, font_name: str, size_pt: float = 10.0) -> None:
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    r = run._element
    rpr = r.get_or_add_rPr()
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), font_name)
    rfonts.set(qn("w:hAnsi"), font_name)
    rfonts.set(qn("w:cs"), font_name)
    rpr.append(rfonts)


def add_code_block(doc: Document, code: str, font_size: float = 10.0) -> None:
    for line in code.strip("\n").split("\n"):
        p = doc.add_paragraph()
        run = p.add_run(line)
        _set_run_font(run, "Courier New", font_size)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.0


def _remove_existing_section(doc: Document) -> None:
    marker = "Дополнение к приложению А: недостающие листинги"
    idx = next(
        (i for i, p in enumerate(doc.paragraphs) if p.text.strip().startswith(marker)),
        None,
    )
    if idx is None:
        return
    start = idx - 1 if idx > 0 else idx
    if idx > 0:
        prev = doc.paragraphs[idx - 1]
        if "w:br" in prev._p.xml and 'w:type="page"' in prev._p.xml:
            start = idx - 1
    body = doc._element.body
    to_remove = [doc.paragraphs[j]._element for j in range(start, len(doc.paragraphs))]
    for el in to_remove:
        body.remove(el)


def main() -> None:
    models_py = (PROJECT / "backend/src/features/database/models.py").read_text(encoding="utf-8")
    chat_py = (PROJECT / "frontend/src/services/chat_service.py").read_text(encoding="utf-8")

    doc = Document(str(DOC_PATH))
    _remove_existing_section(doc)

    doc.add_page_break()

    h = doc.add_paragraph()
    r = h.add_run(
        "Дополнение к приложению А: недостающие листинги и порядок исправления основного файла"
    )
    r.bold = True
    r.font.size = Pt(14)

    doc.add_paragraph()
    p_inst = doc.add_paragraph()
    ri = p_inst.add_run("Инструкция для файла «Курсовая_Вторая.docx»")
    ri.bold = True
    ri.font.size = Pt(12)

    instructions = [
        "Контекст: в приложении А подписи листингов А.1–А.8 должны соответствовать полным "
        "текстам модулей из репозитория ai-job-finder-python. У подписи А.1 должен идти "
        "код models.py, под А.8 — полный chat_service.py. Если после «Листинг А.8» сразу "
        "начинается приложение Б без кода — вставьте листинг А.8 из блока ниже.",
        "Что вставить и куда (минимальное исправление):",
        "1) Откройте «Курсовая_Вторая.docx», найдите строку «Листинг А.8 – Файл "
        "frontend/src/services/chat_service.py». Сразу после строки «Примечание – Источник: "
        "собственная разработка» под этим листингом вставьте полный текст из раздела "
        "«Листинг А.8 (полный текст)» в конце этого файла (шрифт моноширинный, как у "
        "остальных листингов приложения А).",
        "2) Найдите «Листинг А.1 – Файл backend/src/features/database/models.py». Если после "
        "примечания идёт код не models.py (например, начинается с «FastAPI router»), "
        "удалите ошибочный блок кода до следующей подписи «Листинг А.2» и вставьте "
        "полный текст из раздела «Листинг А.1 (полный текст)» ниже.",
        "Если неверно расположены все блоки А.1–А.7 (сдвиг на один модуль): удалите цепочку "
        "семи блоков исходного кода между подписью А.1 и подписью А.8 (оставляя подписи и "
        "примечания на месте). Затем между каждой парой «Листинг А.N» и следующей подписью "
        "вставьте файлы целиком в таком порядке: A.1 — backend/src/features/database/models.py; "
        "A.2 — backend/src/features/search/router.py; A.3 — backend/src/features/search/schemas.py; "
        "A.4 — backend/src/features/llm/service.py; A.5 — backend/src/features/search/service.py; "
        "A.6 — backend/src/features/embedding/save_vectors_service.py; A.7 — frontend/src/app.py. "
        "Листинг А.8 — frontend/src/services/chat_service.py (как в пункте 1). Источники — "
        "папка ai-job-finder-python или блоки в конце файла «Курсовая_Вторая_Демо.docx».",
        "После правок обновите PDF экспортом из Word (Файл → Сохранить как PDF), чтобы "
        "«Курсовая_Вторая.pdf» совпадал с docx.",
    ]
    for para in instructions:
        doc.add_paragraph(para)

    doc.add_paragraph()
    cap1 = doc.add_paragraph()
    c1 = cap1.add_run("Листинг А.1 (полный текст) – Файл backend/src/features/database/models.py")
    c1.bold = True
    c1.font.size = Pt(11)
    add_code_block(doc, models_py, 10.0)
    doc.add_paragraph("Примечание – Источник: собственная разработка")

    doc.add_paragraph()
    cap8 = doc.add_paragraph()
    c8 = cap8.add_run("Листинг А.8 (полный текст) – Файл frontend/src/services/chat_service.py")
    c8.bold = True
    c8.font.size = Pt(11)
    add_code_block(doc, chat_py, 10.0)
    doc.add_paragraph("Примечание – Источник: собственная разработка")

    doc.save(str(DOC_PATH))
    print(f"Updated: {DOC_PATH}")


if __name__ == "__main__":
    main()
