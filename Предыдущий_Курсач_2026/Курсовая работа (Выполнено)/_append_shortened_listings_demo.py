#!/usr/bin/env python3
"""Append shortened chapter listings to Курсовая_Вторая_Демо.docx (end of file only)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parent
DOC_PATH = ROOT / "Курсовая_Вторая_Демо.docx"


def _set_run_font(run, font_name: str, size_pt: float = 10.0) -> None:
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    r = run._element
    rpr = r.get_or_add_rPr()
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), font_name)
    rfonts.set(qn("w:hAnsi"), font_name)
    rfonts.set(qn("w:cs"), font_name)
    rpr.append(rfonts)


def add_code_block(doc: Document, code: str, font_size: float = 9.0) -> None:
    for line in code.strip("\n").split("\n"):
        p = doc.add_paragraph()
        run = p.add_run(line)
        _set_run_font(run, "Courier New", font_size)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.0


LISTINGS: dict[str, str] = {
    "3.1": '''\
from sqlalchemy import String, Text, UniqueConstraint
from sqlalchemy.orm import DeclarativeBase, Mapped, mapped_column
from pgvector.sqlalchemy import Vector

class Base(DeclarativeBase):
    pass

class Vacancy(Base):
    __tablename__ = "vacancies"
    __table_args__ = (UniqueConstraint("url", name="uq_vacancies_url"),)

    id: Mapped[int] = mapped_column(primary_key=True, autoincrement=True)
    title: Mapped[str] = mapped_column(String(512), nullable=False)
    company: Mapped[str] = mapped_column(String(512), nullable=False)
    # … строковые поля salary … work_format (см. полный листинг в приложении А)
    skills: Mapped[str] = mapped_column(Text, nullable=False)
    url: Mapped[str] = mapped_column(String(2048), nullable=False)
    description: Mapped[str] = mapped_column(Text, nullable=False)
    embedding: Mapped[list[float] | None] = mapped_column(Vector(384), nullable=True)''',
    "3.2": '''\
from alembic import op

def upgrade() -> None:
    op.execute("CREATE EXTENSION IF NOT EXISTS vector")
    op.execute(
        """
        CREATE TABLE vacancies (
            id SERIAL PRIMARY KEY,
            title VARCHAR(512) NOT NULL,
            company VARCHAR(512) NOT NULL,
            -- … столбцы salary, payment_frequency, …, description (как в модели Vacancy)
            embedding vector(384),
            CONSTRAINT uq_vacancies_url UNIQUE (url)
        )
        """
    )

def downgrade() -> None:
    op.execute("DROP TABLE IF EXISTS vacancies")''',
    "3.3": '''\
from dataclasses import asdict
from fastapi import APIRouter, HTTPException, Query
from pydantic import BaseModel, HttpUrl
# … импорты констант и ApiErrorResponse

router = APIRouter(prefix=API_V1_PREFIX, tags=[SEARCH_TAG])

class VacancySearchResponse(BaseModel):
    id: int
    title: str
    company: str
    # … остальные поля вакансии
    cosine_distance: float

@router.get(SEARCH_ENDPOINT, response_model=list[VacancySearchResponse], ...)
async def search_vacancies(
    query: str = Query(..., min_length=QUERY_MIN_LENGTH, ...),
    limit: int = Query(default=MAX_VACANCY_RESULTS, ge=1, le=MAX_VACANCY_RESULTS),
) -> list[VacancySearchResponse]:
    normalized_query = query.strip()
    if not normalized_query:
        raise HTTPException(status_code=HTTP_BAD_REQUEST, detail=ERROR_MESSAGE_BLANK_QUERY)
    try:
        search_result = await user_search(normalized_query, limit=limit)
    except RuntimeError as error:
        raise HTTPException(
            status_code=HTTP_INTERNAL_SERVER_ERROR,
            detail=ERROR_MESSAGE_SEARCH_UNAVAILABLE,
        ) from error
    return [
        VacancySearchResponse.model_validate(asdict(vacancy))
        for vacancy in search_result.vacancies
    ]''',
    "3.4": '''\
from dataclasses import dataclass
from features.embedding.schemas import SimilaritySearchResult

@dataclass(frozen=True, slots=True)
class FilterRelaxStep:
    step: int
    filter_key: str
    filter_value: str
    removed_weight: float
    value_count_in_db: int
    candidates_before: int
    candidates_after: int

@dataclass(frozen=True, slots=True)
class UserSearchResult:
    user_query: str
    role_keywords: list[str]
    llm_filters: dict[str, list[dict[str, float | str]]]
    all_value_counts: dict[str, dict[str, int]]
    selected_value_counts: dict[str, dict[str, int]]
    applied_filters: dict[str, list[str]]
    dropped_filters: dict[str, list[str]]
    relax_steps: list[FilterRelaxStep]
    candidate_count: int
    vacancies: list[SimilaritySearchResult]''',
    "3.5": '''\
GROQ_BASE_URL = "https://api.groq.com/openai/v1/chat/completions"
GROQ_MODEL = "openai/gpt-oss-120b"
GROQ_TIMEOUT_SECONDS = 45

VACANCY_FILTER_KEYS = (
    "payment_frequency",
    "experience",
    "employment",
    "hiring_format",
    "schedule",
    "hours",
    "work_format",
)

VACANCY_OPTIONAL_KEYS = ("role_keywords",)

VACANCY_FILTERS_SYSTEM_PROMPT_TEMPLATE = """
Ты — узкоспециализированный ассистент, который переводит запрос кандидата …
в структурированные JSON-фильтры.
ПРАВИЛА:
1. Ответ — только валидный JSON (без markdown).
2. JSON содержит все поля фильтров и role_keywords.
3. Значения — только из списка допустимых (подставляется в шаблон).
… (полный текст промпта и маппинг синонимов — в приложении А)
{allowed_values_json}
""".strip()''',
    "3.6": '''\
# backend/src/common/constants/search.py
DEFAULT_MIN_FILTERED_CANDIDATES = 1
DEFAULT_MAX_VALUES_PER_FILTER = 2
PROTECTED_RELAX_FILTER_KEYS = ("experience", "employment")
PROTECTED_RELAX_WEIGHT_PENALTY = 0.35
ROLE_KEYWORDS_HARD_LOCK = True
DOMAIN_BOOST_STEP = 0.03

# backend/src/features/search/service.py — фрагмент _relax_filters
async def _relax_filters(weighted_filters, *, value_counts, role_keywords, min_candidates):
    active_filters = _clone_non_empty(_plain_filters_from_weighted(weighted_filters))
    weighted_by_key = {k: list(v) for k, v in weighted_filters.items()}
    dropped_filters = _empty_filters()
    relax_steps: list[FilterRelaxStep] = []
    candidate_count = await _count_filtered_candidates(
        active_filters, role_keywords=role_keywords
    )
    if candidate_count >= min_candidates:
        return _with_all_keys(active_filters), dropped_filters, relax_steps, candidate_count

    protected_keys = set(PROTECTED_RELAX_FILTER_KEYS)
    drop_order = _build_field_drop_order(weighted_by_key, protected_keys)
    for step_index, (key, weight) in enumerate(drop_order, start=1):
        removed_values = [value for value, _ in weighted_by_key.get(key, [])]
        if not removed_values:
            continue
        candidates_before = candidate_count
        dropped_filters[key].extend(removed_values)
        active_filters.pop(key, None)
        candidate_count = await _count_filtered_candidates(
            active_filters, role_keywords=role_keywords
        )
        removed_total_count = sum(
            value_counts.get(key, {}).get(value, 0) for value in removed_values
        )
        relax_steps.append(
            FilterRelaxStep(
                step=step_index,
                filter_key=key,
                filter_value=", ".join(removed_values),
                removed_weight=weight,
                value_count_in_db=removed_total_count,
                candidates_before=candidates_before,
                candidates_after=candidate_count,
            )
        )
        if candidate_count >= min_candidates:
            break
    return _with_all_keys(active_filters), dropped_filters, relax_steps, candidate_count''',
    "3.7": '''\
from dto.chat_dto import ChatMessageDto
from dto.vacancy_dto import VacancyDto
from services.backend_api_service import BackendApiService

class ChatService:
    def __init__(self) -> None:
        self.backend_service = BackendApiService()

    @staticmethod
    def _sort_vacancies_by_similarity(vacancies: list[VacancyDto]) -> list[VacancyDto]:
        return sorted(
            vacancies,
            key=lambda v: (
                v.cosine_distance is None,
                v.cosine_distance if v.cosine_distance is not None else float("inf"),
            ),
        )

    def build_response(self, user_prompt: str) -> ChatMessageDto:
        analyzed = self.backend_service.get_vacancy_for_profile(user_prompt)
        if not analyzed.vacancies:
            return ChatMessageDto(role="assistant", content=analyzed.summary)
        sorted_vacancies = self._sort_vacancies_by_similarity(analyzed.vacancies)
        count = len(sorted_vacancies)
        text = f"{analyzed.summary}\\n\\nНайдено вакансий: **{count}**."
        return ChatMessageDto(
            role="assistant", content=text, vacancies=sorted_vacancies
        )''',
    "3.8": '''\
# backend/src/common/constants/embedding.py
EMBEDDING_MODEL_NAME = "paraphrase-multilingual-MiniLM-L12-v2"
EMBEDDING_DIMENSION = 384
DEFAULT_SIMILARITY_TOP_K = 5
VACANCY_EMBED_TITLE_REPEATS = 4
VACANCY_EMBED_STRUCTURE_REPEATS = 3
VACANCY_EMBED_COMPANY_REPEATS = 1
VACANCY_EMBED_SKILLS_REPEATS = 1
VACANCY_EMBED_INCLUDE_DESCRIPTION = True

# save_vectors_service.py — ядро _vacancy_embed_text (без вспомогательных функций)
def _vacancy_embed_text(row) -> str:
    chunks: list[str] = []
    title = _trim(row.title)
    if title:
        job_line = f"Должность: {title}"
        chunks.extend([job_line] * VACANCY_EMBED_TITLE_REPEATS)
    chunks.extend(_repeat_lines("Компания", row.company, VACANCY_EMBED_COMPANY_REPEATS))
    # … для пар (label, attr) из _STRUCTURE_FIELDS — _repeat_lines(..., VACANCY_EMBED_STRUCTURE_REPEATS)
    chunks.extend(_repeat_lines("Зарплата", row.salary, 1))
    chunks.extend(_repeat_lines("Выплаты", row.payment_frequency, 1))
    if not _skip_value(row.skills):
        sk = f"Навыки: {_trim(row.skills)}"
        chunks.extend([sk] * VACANCY_EMBED_SKILLS_REPEATS)
    desc = _description_for_embed(row.description)
    if desc:
        chunks.append(f"Описание: {desc}")
    return "\\n\\n".join(chunks)''',
}


def _remove_existing_shortened_section(doc: Document) -> None:
    """Удаляет ранее добавленный блок «Исправленные сокращённые листинги…» до конца документа."""
    marker = "Исправленные сокращённые листинги для текста глав"
    idx = next(
        (i for i, p in enumerate(doc.paragraphs) if p.text.strip().startswith(marker)),
        None,
    )
    if idx is None:
        return
    start = idx
    if idx > 0:
        prev = doc.paragraphs[idx - 1]
        if "w:br" in prev._p.xml and 'w:type="page"' in prev._p.xml:
            start = idx - 1
    body = doc._element.body
    to_remove = [doc.paragraphs[j]._element for j in range(start, len(doc.paragraphs))]
    for el in to_remove:
        body.remove(el)


def main() -> None:
    doc = Document(str(DOC_PATH))
    _remove_existing_shortened_section(doc)

    doc.add_page_break()
    p0 = doc.add_paragraph()
    r0 = p0.add_run(
        "Исправленные сокращённые листинги для текста глав (фрагменты)"
    )
    r0.bold = True
    r0.font.size = Pt(14)

    doc.add_paragraph(
        "Ниже — варианты листингов с акцентом на ключевые фрагменты. Полные тексты модулей "
        "см. в приложении А и в репозитории ai-job-finder-python. Нумерация соответствует "
        "листингам 3.1–3.8 основного текста."
    )

    for num in ["3.1", "3.2", "3.3", "3.4", "3.5", "3.6", "3.7", "3.8"]:
        doc.add_paragraph()
        ph = doc.add_paragraph()
        rh = ph.add_run(f"Листинг {num}")
        rh.bold = True
        rh.font.size = Pt(11)
        add_code_block(doc, LISTINGS[num])

    doc.add_paragraph()
    pn = doc.add_paragraph()
    rn = pn.add_run("Примечание — Источник: собственная разработка")
    rn.italic = True
    rn.font.size = Pt(10)

    doc.save(str(DOC_PATH))
    print(f"Appended shortened listings to: {DOC_PATH}")


if __name__ == "__main__":
    main()
